# -*- coding: utf-8 -*-
"""专家评审结果统计：检查器裁定（precision / 估计 recall / κ）与终止式裁定（写回 YAML）。

检查器裁定
----------
读 ``docs/expert_review/检查器专家裁定表_评审人*.xlsx`` 与答案键
``runs/private/expert_review_key.json``（由 ``expert_pack`` 生成）。

* **precision**（按抽样层）：检查器报错拍中，评审判「有错」的比例；「不确定」不计入分母。
  另报「类型一致」的 precision：判有错且所选错误类型映射到该层规则。
* **漏报率**：未报错拍中被判「有错」的比例；
* **估计 recall**：按抽样框总体加权，TP̂ = Σ_层 precision_层 × 总体报错数_层，
  FN̂ = 漏报率 × 总体未报错数，recall = TP̂ / (TP̂ + FN̂)。层按条目的抽样层计，
  同一拍报多条规则的只计一次（按其抽样层），因此是近似。
* **κ**：两位评审在「有错/无错/不确定」三类与去掉不确定后的二类上的 Cohen's κ。
* 置信区间：比例用 Wilson 95%。

终止式裁定
----------
读 ``docs/cadence_review/终止式评审表*.xlsx`` 的「② 终止式类型判定」。多份表结论一致
（或只有一份）才写回；不一致的列为待仲裁。``--apply`` 时把结论写入
``grammar.yaml``：``type`` 改为对应类型，并记 ``adjudication`` 字段；判 NOT_CADENCE 的
把 ``p`` 置 0（生成器不再选用，检查器的终止式命中也不再计入）。默认只报告不写。

用法::

    venv/bin/python -m score_learning.grammar.expert_score            # 只统计
    venv/bin/python -m score_learning.grammar.expert_score --apply    # 并写回 YAML
"""
from __future__ import annotations

import argparse
import glob
import json
import math
import os
import re
import sys
from collections import Counter, defaultdict
from typing import Any, Dict, List, Optional, Tuple

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from score_learning.grammar.expert_pack import (  # noqa: E402
    KEY_PATH, OUT_DIR, SAMPLED_RULES, TYPE_TO_RULE,
)

CAD_DIR = os.path.join(ROOT, 'docs', 'cadence_review')
YAML = os.path.join(ROOT, 'skills', 'composition', 'genre', 'packs',
                    'sacred', 'chorale', 'grammar.yaml')
OUT = os.path.join(ROOT, 'runs', 'experiments')

VERDICT_TO_TYPE = {
    'KEEP_OTHER': '其他终止', 'RECODE_PAC': '完全正格终止 PAC',
    'RECODE_IAC': '不完全正格终止 IAC', 'RECODE_PC': '变格终止 PC',
    'RECODE_HC': '半终止 HC', 'RECODE_DC': '阻碍终止 DC',
    'RECODE_SECONDARY': '离调终止（副属→目标）', 'RECODE_TONIC': '主和弦终止',
    'MODULATED': '转调终止（专家裁定）', 'NOT_CADENCE': '非终止（专家裁定）',
}


# ------------------------------------------------------------------ 统计工具
def wilson(k: int, n: int, z: float = 1.96) -> Tuple[Optional[float], Optional[float]]:
    if n == 0:
        return None, None
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return round(max(0.0, c - h), 4), round(min(1.0, c + h), 4)


def cohen_kappa(a: List[str], b: List[str]) -> Optional[float]:
    pairs = [(x, y) for x, y in zip(a, b) if x and y]
    n = len(pairs)
    if n == 0:
        return None
    po = sum(x == y for x, y in pairs) / n
    ca, cb = Counter(x for x, _ in pairs), Counter(y for _, y in pairs)
    pe = sum(ca[k] * cb.get(k, 0) for k in ca) / (n * n)
    return round((po - pe) / (1 - pe), 4) if pe < 1 else 1.0


# ------------------------------------------------------------------ 检查器裁定
def read_checker_sheets(paths: Optional[List[str]] = None) -> Dict[str, Dict[str, Dict[str, str]]]:
    """{评审人: {条目: {'verdict', 'type', 'label', 'note'}}}；空表返回空 dict。"""
    from openpyxl import load_workbook
    paths = paths or sorted(glob.glob(os.path.join(OUT_DIR, '检查器专家裁定表_评审人*.xlsx')))
    out = {}
    for p in paths:
        rater = os.path.splitext(os.path.basename(p))[0].split('评审人')[-1]
        ws = load_workbook(p, data_only=True)['评审表']
        rows = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or not r[0]:
                continue
            rows[str(r[0])] = {'verdict': (r[3] or '').strip(), 'type': (r[4] or '').strip(),
                               'label': (r[5] or '').strip(), 'note': (r[6] or '')}
        if any(v['verdict'] for v in rows.values()):
            out[rater] = rows
    return out


def score_checker(key: Dict[str, Any], sheets: Dict[str, Dict[str, Dict[str, str]]]
                  ) -> Dict[str, Any]:
    items = {it['item']: it for it in key['items']}
    pop = key['population']
    res: Dict[str, Any] = {'raters': sorted(sheets), 'n_items': len(items), 'per_rater': {}}
    for rater, rows in sheets.items():
        strata: Dict[str, Counter] = defaultdict(Counter)
        for iid, it in items.items():
            v = rows.get(iid, {}).get('verdict', '')
            if v not in ('有错', '无错'):
                strata[it['stratum']]['unsure_or_blank'] += 1
                continue
            c = strata[it['stratum']]
            c['n'] += 1
            c['yes'] += v == '有错'
            if it['stratum'] != 'clean' and v == '有错':
                c['type_match'] += TYPE_TO_RULE.get(rows[iid]['type']) == it['stratum'] or (
                    it['stratum'] == 'R03' and rows[iid]['type'] in ('声部交叉', '声部超越'))
            lab = rows.get(iid, {}).get('label', '')
            c['label_judged'] += lab in ('正确', '有误')
            c['label_ok'] += lab == '正确'
        per = {}
        for s, c in strata.items():
            lo, hi = wilson(c['yes'], c['n'])
            per[s] = {'n': c['n'], 'judged_violation': c['yes'],
                      'rate': round(c['yes'] / c['n'], 4) if c['n'] else None,
                      'ci95': [lo, hi], 'type_match': c['type_match'],
                      'unsure_or_blank': c['unsure_or_blank']}
        # 总体加权的 precision 与 recall（两来源合并）
        tp_hat = 0.0
        flagged_pop = 0
        for r in SAMPLED_RULES:
            n_r = sum(pop[s]['by_rule'][r] for s in pop)
            if per.get(r, {}).get('rate') is not None:
                tp_hat += per[r]['rate'] * n_r
                flagged_pop += n_r
        clean_pop = sum(pop[s]['clean'] for s in pop)
        miss = per.get('clean', {}).get('rate')
        fn_hat = (miss or 0.0) * clean_pop
        label_j = sum(c['label_judged'] for c in strata.values())
        label_ok = sum(c['label_ok'] for c in strata.values())
        res['per_rater'][rater] = {
            'per_stratum': per,
            'precision_weighted': round(tp_hat / flagged_pop, 4) if flagged_pop else None,
            'miss_rate_unflagged': miss,
            'recall_estimated': round(tp_hat / (tp_hat + fn_hat), 4) if tp_hat + fn_hat else None,
            'chord_label_correct': round(label_ok / label_j, 4) if label_j else None,
            'chord_label_judged': label_j,
        }
    if len(sheets) >= 2:
        ra, rb = sorted(sheets)[:2]
        ids = sorted(items)
        va = [sheets[ra].get(i, {}).get('verdict', '') for i in ids]
        vb = [sheets[rb].get(i, {}).get('verdict', '') for i in ids]
        bin_a = [x if x in ('有错', '无错') else '' for x in va]
        bin_b = [x if x in ('有错', '无错') else '' for x in vb]
        res['kappa'] = {'raters': [ra, rb], 'three_way': cohen_kappa(va, vb),
                        'binary': cohen_kappa(bin_a, bin_b),
                        'n_both_binary': sum(1 for x, y in zip(bin_a, bin_b) if x and y)}
    return res


# ------------------------------------------------------------------ 终止式裁定
def read_cadence_sheets(paths: Optional[List[str]] = None) -> Dict[str, Dict[str, Dict[str, str]]]:
    from openpyxl import load_workbook
    paths = paths or sorted(glob.glob(os.path.join(CAD_DIR, '终止式评审表*.xlsx')))
    out = {}
    for p in paths:
        ws = load_workbook(p, data_only=True)['评审表']
        rows = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or not r[0]:
                continue
            rows[str(r[0])] = {'real': r[6] or '', 'verdict': (r[7] or '').strip(),
                               'confidence': r[11] or '', 'basis': r[12] or '',
                               'note': r[13] or ''}
        if any(v['verdict'] for v in rows.values()):
            out[os.path.basename(p)] = rows
    return out


def resolve_cadences(sheets: Dict[str, Dict[str, Dict[str, str]]]) -> Dict[str, Any]:
    ids = sorted({c for s in sheets.values() for c in s}, key=lambda x: int(x[1:]))
    agreed, conflict, pending = {}, {}, []
    for cid in ids:
        vs = {name: s[cid]['verdict'] for name, s in sheets.items()
              if cid in s and s[cid]['verdict']}
        if not vs:
            pending.append(cid)
        elif len(set(vs.values())) == 1:
            v = next(iter(vs.values()))
            first = next(s[cid] for s in sheets.values() if cid in s and s[cid]['verdict'])
            agreed[cid] = {'verdict': v, 'confidence': first['confidence'],
                           'basis': first['basis'], 'sheets': sorted(vs)}
        else:
            conflict[cid] = vs
    return {'agreed': agreed, 'conflict': conflict, 'pending': pending,
            'verdict_hist': dict(Counter(v['verdict'] for v in agreed.values()))}


def apply_cadences(agreed: Dict[str, Dict[str, Any]], path: str = YAML) -> List[str]:
    """把一致的裁定写回 YAML；返回改动过的 C 编号。UNDETERMINED 不改。"""
    from score_learning.grammar import script as gs
    d = gs.load(path)
    changed = []
    for mode, targets in (d.get('cadences') or {}).items():
        for _tgt, blk in targets.items():
            for f in blk.get('forms') or []:
                a = agreed.get(f['id'])
                if not a or a['verdict'] not in VERDICT_TO_TYPE:
                    continue
                f['type'] = VERDICT_TO_TYPE[a['verdict']]
                if f.get('say'):
                    f['say'] = re.sub(r'，[^，）]*）$', f'，{f["type"]}）', f['say'])
                f['adjudication'] = {'verdict': a['verdict'], 'confidence': a['confidence'],
                                     'basis': a['basis'], 'by': 'expert'}
                if a['verdict'] == 'NOT_CADENCE':
                    f['p'] = 0.0
                changed.append(f['id'])
    gs.save(d, path, os.path.splitext(path)[0] + '.md')
    return changed


def refined_agreement(key: Dict[str, Any], sheets: Dict[str, Dict[str, Dict[str, str]]]
                      ) -> Dict[str, Any]:
    """同一批条目，按 refined 检查器重新判，与评审结论对照。

    抽样以原检查器分层，故按设计权重回推总体：报错层每条代表 flagged/145 拍，
    未报错层每条代表 clean/50 拍。
    """
    from score_learning.grammar.expert_pack import collect
    cands, _ = collect(key['seed'], refined=True)
    new_rules = {(c['source'], c['piece'], c['k']): c['rules'] for c in cands}
    pop = key['population']
    n_flag = sum(1 for it in key['items'] if it['stratum'] != 'clean')
    n_clean = len(key['items']) - n_flag
    w_flag = sum(p['flagged'] for p in pop.values()) / n_flag
    w_clean = sum(p['clean'] for p in pop.values()) / n_clean
    out: Dict[str, Any] = {'weights': {'flagged': round(w_flag, 3), 'clean': round(w_clean, 3)},
                           'items_changed': [], 'per_rater': {}}
    for it in key['items']:
        nr = new_rules.get((it['source'], it['piece'], it['k']), [])
        if sorted(nr) != sorted(it['rules']):
            out['items_changed'].append({'item': it['item'], 'old': it['rules'], 'new': nr})
    for rater, rows in sheets.items():
        res = {}
        for arm in ('original', 'refined'):
            c = Counter()
            for it in key['items']:
                v = rows.get(it['item'], {}).get('verdict', '')
                if v not in ('有错', '无错'):
                    continue
                flag = bool(it['rules'] if arm == 'original'
                            else new_rules.get((it['source'], it['piece'], it['k'])))
                w = w_flag if it['stratum'] != 'clean' else w_clean
                c[(flag, v == '有错')] += w
            tp, fp, fn = c[(True, True)], c[(True, False)], c[(False, True)]
            res[arm] = {'precision': round(tp / (tp + fp), 4) if tp + fp else None,
                        'recall': round(tp / (tp + fn), 4) if tp + fn else None,
                        'flagged_judged': sum(1 for it in key['items']
                                              if (it['rules'] if arm == 'original' else
                                                  new_rules.get((it['source'], it['piece'], it['k'])))
                                              and rows.get(it['item'], {}).get('verdict') in ('有错', '无错'))}
        removed = [x for x in out['items_changed'] if x['old'] and not x['new']]
        res['removed_flags'] = len(removed)
        res['removed_judged_no_error'] = sum(
            1 for x in removed if rows.get(x['item'], {}).get('verdict') == '无错')
        added = [x for x in out['items_changed'] if not x['old'] and x['new']]
        res['added_flags'] = len(added)
        res['added_judged_error'] = sum(
            1 for x in added if rows.get(x['item'], {}).get('verdict') == '有错')
        out['per_rater'][rater] = res
    return out


# ------------------------------------------------------------------ 主流程
def run(apply: bool = False, save: bool = True, refined: bool = True) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    if os.path.isfile(KEY_PATH):
        key = json.load(open(KEY_PATH, encoding='utf-8'))
        sheets = read_checker_sheets()
        out['checker'] = (score_checker(key, sheets) if sheets
                          else {'status': 'awaiting rater sheets', 'n_items': len(key['items'])})
        if sheets and refined:
            out['checker_refined'] = refined_agreement(key, sheets)
    else:
        out['checker'] = {'status': 'answer key missing; run expert_pack first'}
    cs = read_cadence_sheets()
    if cs:
        res = resolve_cadences(cs)
        if apply and res['agreed']:
            res['applied'] = apply_cadences(res['agreed'])
        out['cadence'] = res
    else:
        out['cadence'] = {'status': 'awaiting cadence verdicts'}
    if save:
        os.makedirs(OUT, exist_ok=True)
        with open(os.path.join(OUT, 'e9_expert_adjudication.json'), 'w', encoding='utf-8') as f:
            json.dump(out, f, ensure_ascii=False, indent=1)
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description='专家评审结果统计')
    ap.add_argument('--apply', action='store_true', help='把一致的终止式裁定写回 grammar.yaml')
    ap.add_argument('--no-refined', action='store_true',
                    help='不做按现行（refined）检查器重判同一批条目的对照')
    a = ap.parse_args(argv)
    out = run(apply=a.apply, refined=not a.no_refined)
    print(json.dumps(out, ensure_ascii=False, indent=1)[:4000])
    print('产物 → runs/experiments/e9_expert_adjudication.json')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
